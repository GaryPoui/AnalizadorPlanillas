from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Manual_de_Administracion_y_Operacion_PriceBot.docx"

DARK_BLUE = "183B56"
PALE_BLUE = "EAF3F8"
LIGHT_GRAY = "F4F6F8"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_table_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


def add_number(doc, text):
    if not hasattr(add_number, "counter"):
        add_number.counter = 0
    add_number.counter += 1
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.28)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(f"{add_number.counter}. {text}")
    return paragraph


def reset_numbering():
    add_number.counter = 0


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    keep_table_row_together(header)
    for index, title in enumerate(headers):
        cell = header.cells[index]
        cell.text = title
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        keep_table_row_together(table.rows[-1])
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = BLACK
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08

for style_name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Heading 3" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(7)

for list_style in ("List Bullet", "List Bullet 2"):
    styles[list_style].paragraph_format.space_after = Pt(2)
    styles[list_style].font.size = Pt(10)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Manual de administración y operación de PriceBot")
subtitle = doc.add_paragraph()
subtitle.add_run("Guía de puesta en marcha seguridad usuarios tokens créditos y uso diario").bold = True
subtitle.paragraph_format.space_after = Pt(18)

metadata = add_table(
    doc,
    ["Dato", "Valor"],
    [
        ["Destinatario", "Área de Administración y soporte técnico"],
        ["Sistema", "AnalizadorPlanillas PriceBot"],
        ["Servidor de referencia", "192.168.190.146 en la red 192.168.190.0/24"],
        ["Versión del manual", "6 de octubre de 2026"],
        ["Repositorio", "github.com/GaryPoui/AnalizadorPlanillas"],
    ],
    [1.65, 5.05],
)

doc.add_heading("Objetivo y alcance", level=1)
doc.add_paragraph(
    "Este manual permite que el área de Administración opere PriceBot, gestione accesos, "
    "mantenga las cuentas externas y controle el consumo de créditos sin depender del equipo "
    "que desarrolló la aplicación. También establece qué datos pueden compartirse y qué secretos "
    "deben permanecer únicamente en la computadora servidor."
)
doc.add_paragraph(
    "PriceBot recibe listas de precios en PDF, Excel, CSV o imagen y genera una planilla normalizada. "
    "La extracción local no requiere créditos. El uso de IA externa es opcional y solo debe activarse "
    "cuando exista autorización para enviar contenido del documento a Anthropic."
)

doc.add_heading("Estado operativo y decisión de seguridad", level=1)
doc.add_paragraph(
    "El lanzador admite administración local por HTTP en localhost, pero exige HTTPS para "
    "conexiones desde otras computadoras. Para habilitar el acceso LAN se debe instalar un "
    "certificado confiable cuya SAN incluya la IP/nombre publicado y configurar TLS en frontend "
    "y API. HTTP no cifra credenciales ni documentos y no debe habilitarse para uso remoto."
)
add_bullet(doc, "No exponer los puertos 3000 u 8000 directamente a Internet.")
add_bullet(doc, "Permitir acceso únicamente desde la subred autorizada y el perfil de red Privado.")
add_bullet(doc, "Mantener PRICEBOT_ALLOW_EXTERNAL_AI en 0 si el documento no puede salir de la organización.")
add_bullet(doc, "No enviar contraseñas, claves API ni contraseñas de aplicación por correo o chat.")

doc.add_heading("Requisitos para habilitar usuarios remotos", level=1)
doc.add_paragraph(
    "La configuración actual requiere dos acciones de TI antes de habilitar la operación desde otras PCs: "
    "definir el secreto de sesión y configurar un certificado TLS confiable. Sin secreto, los usuarios "
    "remotos reciben un error de autenticación; sin TLS, el servidor bloquea HTTP remoto por seguridad."
)
reset_numbering()
add_number(doc, "En pricebot\\.env, establecer PRICEBOT_SESSION_SECRET con una clave aleatoria nueva. En PowerShell generar bytes con RNG criptográfico; copiar el resultado directamente al .env y no compartirlo:")
secret_example = doc.add_paragraph()
secret_example.paragraph_format.left_indent = Inches(0.3)
secret_run = secret_example.add_run(
    "$bytes = New-Object byte[] 48\n"
    "$rng = [System.Security.Cryptography.RandomNumberGenerator]::Create()\n"
    "$rng.GetBytes($bytes)\n"
    "[Convert]::ToBase64String($bytes)\n"
    "$rng.Dispose()"
)
secret_run.font.name = "Consolas"
secret_run.font.size = Pt(9)
add_number(doc, "Guardar .env como UTF-8 sin BOM y reiniciar el servidor. No registrar el secreto en este manual, correos, chats ni Git.")
add_number(doc, "TI debe proveer un certificado TLS confiable para el nombre/IP del servidor, con esa identidad en SAN, y configurar sus rutas antes de ejecutar iniciar-servidor.bat.")
tls_example = doc.add_paragraph()
tls_example.paragraph_format.left_indent = Inches(0.3)
tls_run = tls_example.add_run(
    '$env:PRICEBOT_TLS_CERTFILE="C:\\ruta\\pricebot-cert.pem"\n'
    '$env:PRICEBOT_TLS_KEYFILE="C:\\ruta\\pricebot-key.pem"\n'
    ".\\iniciar-servidor.bat"
)
tls_run.font.name = "Consolas"
tls_run.font.size = Pt(9)
add_bullet(doc, "La clave privada TLS también es secreta. No copiarla al repositorio ni reenviarla; limitar permisos NTFS.")

links_heading = doc.add_heading("Enlaces y puntos de acceso", level=1)
links_heading.paragraph_format.page_break_before = True
links = [
    ("Administración local", "http://127.0.0.1:3000", "Solo desde la PC servidor; acceso como Administrador local"),
    ("Acceso de usuarios en LAN", "https://192.168.190.146:3000", "Solo después de configurar y confiar el certificado TLS"),
    ("Estado de la API", "https://192.168.190.146:8000/health", "Diagnóstico remoto HTTPS; debe responder status ok"),
    ("Repositorio del proyecto", "https://github.com/GaryPoui/AnalizadorPlanillas", "Código, historial y actualizaciones"),
    ("Claves de Claude Platform", "https://platform.claude.com/settings/keys", "Crear y rotar ANTHROPIC_API_KEY"),
    ("Facturación de Claude Platform", "https://platform.claude.com/settings/billing", "Comprar créditos y configurar recarga"),
    ("Uso de Claude Platform", "https://platform.claude.com/usage", "Controlar tokens y costos"),
    ("Estado de Anthropic", "https://status.anthropic.com", "Ver incidentes del proveedor"),
    ("Correo Gmail administrativo", "https://mail.google.com/", "Cuenta responsable: compras@dynamicenergy.com.ar"),
    ("Seguridad Google", "https://myaccount.google.com/security", "Activar segundo factor y crear contraseña de aplicación"),
    ("Contraseñas de aplicación Google", "https://myaccount.google.com/apppasswords", "Rotar la credencial SMTP de Gmail"),
]
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
set_table_borders(table)
for idx, text in enumerate(("Recurso", "Enlace", "Uso")):
    cell = table.rows[0].cells[idx]
    cell.text = text
    set_cell_shading(cell, DARK_BLUE)
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
repeat_table_header(table.rows[0])
keep_table_row_together(table.rows[0])
for row_index, (name, url, use) in enumerate(links):
    cells = table.add_row().cells
    keep_table_row_together(table.rows[-1])
    cells[0].text = name
    add_hyperlink(cells[1].paragraphs[0], url, url)
    cells[2].text = use
    for cell in cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if row_index % 2:
            set_cell_shading(cell, PALE_BLUE)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(9)
doc.add_paragraph()

doc.add_heading("Inventario de cuentas y responsables", level=1)
add_table(
    doc,
    ["Cuenta o rol", "Identificador", "Responsabilidad", "Secreto documentado"],
    [
        ["Administrador local de PriceBot", "Administrador local", "Opera usuarios desde localhost en el servidor", "No aplica"],
        ["Cuenta Windows del servidor", "Responsable TI a confirmar", "Iniciar, actualizar y respaldar la aplicación", "No"],
        ["Gmail administrativo", "compras@dynamicenergy.com.ar", "Confirmaciones de acceso y alertas de consumo", "No; usar contraseña de aplicación"],
        ["Usuarios finales", "Correos visibles en la pestaña Usuarios", "Ingresar desde otras PCs", "No; las contraseñas no son recuperables"],
        ["Claude Platform", "Titular a registrar por Administración", "Claves API, créditos, facturación y uso", "No"],
        ["GitHub", "Repositorio GaryPoui AnalizadorPlanillas", "Actualizar el código y conservar historial", "No"],
    ],
    [1.35, 1.8, 2.55, 1.0],
)
doc.add_paragraph(
    "Administración/TI debe confirmar titulares y recuperación de Claude Platform y GitHub. "
    "Nunca se incluyen contraseñas en este manual. Guardar secretos en el gestor corporativo o "
    "en pricebot\\.env con permisos NTFS restringidos. La administración de usuarios de PriceBot "
    "se realiza localmente en el servidor; no existe una contraseña maestra compartida."
)

doc.add_heading("Inicio y detención del servidor", level=1)
reset_numbering()
add_number(doc, "Iniciar sesión en Windows en la computadora servidor.")
add_number(doc, "Abrir la carpeta de instalación de AnalizadorPlanillas indicada por TI.")
add_number(doc, "Hacer doble clic en iniciar-servidor.bat y mantener abiertas las ventanas del backend y frontend.")
add_number(doc, "Comprobar http://127.0.0.1:3000 en el servidor. Para LAN, configurar TLS y usar https://192.168.190.146:3000.")
add_number(doc, "Verificar que /health responda status ok antes de habilitar el uso.")
doc.add_paragraph(
    "Para detener el servicio, cerrar las ventanas iniciadas por PriceBot. No apagar la PC mientras "
    "haya extracciones activas. Después de un reinicio de Windows, ejecutar nuevamente el lanzador."
)

doc.add_heading("Primera puesta en marcha del administrador", level=1)
reset_numbering()
add_number(doc, "TI instala Python y las dependencias ejecutando instalar.bat. Crear pricebot\\.env desde pricebot\\.env.example si aún no existe; nunca reemplazar un .env existente sin respaldarlo.")
add_number(doc, "Acordar con Administración si los documentos pueden enviarse a Anthropic. Mantener PRICEBOT_ALLOW_EXTERNAL_AI=0 si no hay aprobación; solo entonces cargar ANTHROPIC_API_KEY y cambiarlo a 1.")
add_number(doc, "Cargar la cuenta Gmail autorizada, su contraseña de aplicación y PRICEBOT_ALERT_EMAIL siguiendo la sección Gmail. Definir el presupuesto mensual aprobado para habilitar alertas de umbral.")
add_number(doc, "Generar PRICEBOT_SESSION_SECRET en el servidor, guardarlo solo en .env y restringir permisos NTFS. Aplicar HTTPS confiable en frontend/API antes de permitir acceso desde otras PCs.")
add_number(doc, "Revisar PRICEBOT_PUBLIC_URL y PRICEBOT_ALLOWED_ORIGINS para que usen el mismo HTTPS y nombre/IP cubierto por el certificado. Guardar .env como UTF-8 sin BOM.")
add_number(doc, "Ejecutar iniciar-servidor.bat. Abrir http://127.0.0.1:3000 en el servidor y comprobar el estado del servicio. El administrador es un acceso local de Windows, no una cuenta web compartida.")
add_number(doc, "En la pestaña Usuarios, crear el primer usuario con su correo y contraseña inicial (mínimo 10 caracteres). Comunicar esa contraseña por un canal seguro distinto del correo de confirmación.")
add_number(doc, "Pedir al usuario que confirme el correo dentro de las 24 horas. Verificar que figure Confirmado y probar el acceso desde el equipo autorizado.")
doc.add_paragraph(
    "Si todavía no hay usuarios activos, el modo de alta inicial solo está disponible desde localhost. "
    "Con el lanzador oficial, PRICEBOT_LOCAL_ADMIN_BYPASS mantiene la pestaña de administración en el servidor; "
    "las operaciones de usuarios siguen rechazando solicitudes remotas. No habilitar ese bypass en un proxy público."
)

doc.add_heading("Recorrido del usuario: invitación a descarga", level=1)
reset_numbering()
add_number(doc, "Abrir el mensaje de invitación enviado a su correo y entrar al enlace antes de que venza (24 horas por defecto). Debe usar la dirección HTTPS confiable indicada por TI.")
add_number(doc, "Abrir PriceBot, escribir el correo registrado y la contraseña inicial entregada por el administrador y seleccionar Ingresar. Si el correo cambió o no puede entrar, solicitar al administrador que actualice o restablezca la cuenta.")
add_number(doc, "En Subir Archivos, hacer clic en la zona de carga o arrastrar archivos. Formatos admitidos: PDF, XLS, XLSX, XLSM, DOCX, CSV, JPG, JPEG, PNG y WEBP; máximo 50 MB por archivo.")
add_number(doc, "Opcionalmente completar el CUIT del proveedor. Revisar la lista de archivos y quitar cualquier archivo equivocado antes de iniciar.")
add_number(doc, "Seleccionar Extraer con IA. XLS/XLSX/XLSM/CSV se procesan localmente; PDF y DOCX usan extracción de texto/tablas y solo llaman a Anthropic si está habilitado; las imágenes requieren IA externa.")
add_number(doc, "Esperar el estado Completado. Si aparece un error, revisar el mensaje y pedir ayuda al administrador; no volver a subir documentos confidenciales a servicios externos.")
add_number(doc, "En Resultados, revisar cantidad de filas, calidad, advertencias, precios vacíos y filas enviadas a revisión. Comparar códigos y precios con el original, especialmente cuando la calidad sea baja o el documento esté escaneado.")
add_number(doc, "Cuando el resultado sea aceptable, elegir Descargar XLS o Descargar XLSX y guardar la planilla en la ubicación corporativa autorizada. Validar una muestra antes de importarla al sistema de compras.")
add_number(doc, "Al terminar, cerrar sesión y borrar de la lista los archivos que ya no necesite procesar. PriceBot no conserva el historial de resultados por defecto.")

doc.add_heading("Uso diario y soporte", level=1)
reset_numbering()
add_number(doc, "Administración inicia PriceBot y comprueba /health antes del comienzo de la jornada.")
add_number(doc, "Usuarios reportan fallas con hora aproximada y tipo de archivo, sin adjuntar listas confidenciales.")
add_number(doc, "Administración revisa Gmail, Billing/Usage de Anthropic y costs_log.jsonl; no usar el costo estimado local como saldo contable.")
add_number(doc, "Al cerrar la jornada, asegurar respaldos cifrados de .env y users.db según la política de TI.")

doc.add_heading("Formatos y tratamiento de datos", level=2)
add_table(
    doc,
    ["Formato", "Procesamiento", "IA externa", "Control principal"],
    [
        ["XLS XLSX XLSM CSV", "Local con pandas y librerías de Excel", "No por defecto", "Firma, estructura ZIP, tamaño y fórmulas de salida"],
        ["PDF", "Texto tablas y reglas locales; complemento opcional", "Solo si está autorizada", "Firma, tamaño, páginas y tiempo máximo"],
        ["DOCX", "Texto y tablas del documento", "Solo si está autorizada y hace falta", "Estructura Office ZIP y límites de expansión"],
        ["JPG JPEG PNG WEBP", "Lectura de imagen y visión", "Necesaria", "Firma, integridad y límite de píxeles"],
    ],
    [1.15, 2.15, 1.35, 2.05],
)

doc.add_heading("Administración de usuarios", level=1)
doc.add_paragraph(
    "Abrir http://127.0.0.1:3000 en la PC servidor y entrar en la pestaña Usuarios. Esta pestaña y "
    "sus endpoints solo aceptan solicitudes cuyo origen real y encabezado Host sean de loopback."
)
doc.add_heading("Crear un usuario", level=2)
reset_numbering()
add_number(doc, "Ingresar el correo y una contraseña inicial de al menos 10 caracteres.")
add_number(doc, "Seleccionar Crear y enviar confirmación.")
add_number(doc, "Pedir al usuario que abra el enlace recibido dentro de la LAN.")
add_number(doc, "Comprobar que el estado cambie de Pendiente a Confirmado.")
doc.add_heading("Modificar un usuario", level=2)
reset_numbering()
add_number(doc, "Seleccionar Modificar junto al usuario.")
add_number(doc, "Cambiar el correo, la contraseña o ambos. Dejar la contraseña vacía para conservarla.")
add_number(doc, "Guardar. Un cambio de correo genera una nueva confirmación; un cambio de contraseña cierra las sesiones anteriores.")
doc.add_heading("Deshabilitar o eliminar un usuario", level=2)
add_bullet(doc, "Deshabilitar conserva el registro y corta el acceso. Usar esta opción para ausencias o bloqueos temporales.")
add_bullet(doc, "Eliminar borra definitivamente el registro y revoca el acceso. Usarlo cuando la persona deja de estar autorizada.")
add_bullet(doc, "Reenviar solo funciona para usuarios pendientes y habilitados.")

doc.add_heading("Gestión de tokens cuentas y créditos", level=1)
doc.add_heading("Clave de Anthropic", level=2)
reset_numbering()
add_number(doc, "Ingresar a Claude Platform con una cuenta organizacional y doble factor.")
add_number(doc, "Abrir Settings y API keys. Crear una clave identificada como PriceBot servidor.")
add_number(doc, "Copiarla una sola vez en ANTHROPIC_API_KEY dentro de pricebot\\.env.")
add_number(doc, "Establecer PRICEBOT_ALLOW_EXTERNAL_AI=1 únicamente con autorización de tratamiento de datos.")
add_number(doc, "Reiniciar PriceBot y ejecutar una prueba no confidencial.")
doc.add_paragraph(
    "Rotación recomendada: crear una clave nueva, reemplazarla en el servidor, reiniciar, probar y recién "
    "entonces revocar la anterior. Si una clave aparece en un chat, correo, captura o repositorio, revocarla inmediatamente."
)
doc.add_heading("Compra y control de créditos", level=2)
reset_numbering()
add_number(doc, "Abrir la página Billing de Claude Platform y revisar la organización seleccionada.")
add_number(doc, "Seleccionar Buy credits, definir el importe y confirmar el medio de pago autorizado.")
add_number(doc, "Configurar auto reload solo si Administración aprueba el umbral y el monto de recarga.")
add_number(doc, "Revisar Usage por modelo y API key, y contrastarlo con costs_log.jsonl del servidor.")
doc.add_paragraph(
    "Anthropic informa que la API usa créditos prepagados, que las llamadas exitosas consumen saldo y "
    "que los créditos comprados vencen un año después y no son reembolsables. Los roles Developer, "
    "Billing y Admin pueden consultar uso y costos; asignar el mínimo rol necesario."
)
doc.add_heading("Correo Gmail para confirmaciones", level=2)
reset_numbering()
add_number(doc, "Confirmar que compras@dynamicenergy.com.ar es la cuenta autorizada y mantener activa su verificación en dos pasos.")
add_number(doc, "Crear una contraseña de aplicación exclusiva para PriceBot desde Seguridad de Google.")
add_number(doc, "En pricebot\\.env, configurar SMTP_HOST=smtp.gmail.com, puerto 587, usuario y remitente compras@dynamicenergy.com.ar; habilitar STARTTLS y desactivar SSL.")
add_number(doc, "Pegar la contraseña de aplicación sin espacios en PRICEBOT_SMTP_PASSWORD. No usar la contraseña normal de Gmail ni incluirla en este manual.")
add_number(doc, "Configurar PRICEBOT_ALERT_EMAIL=compras@dynamicenergy.com.ar, guardar .env como UTF-8 sin BOM y reiniciar PriceBot.")
add_number(doc, "Probar el SMTP creando un usuario temporal, confirmar recepción y luego eliminarlo.")
doc.add_paragraph("Si Google Workspace no ofrece contraseñas de aplicación, solicitar a TI un relay SMTP autorizado; no reducir la seguridad de la cuenta.")

doc.add_heading("Alertas y carga de créditos", level=2)
reset_numbering()
add_number(doc, "Definir PRICEBOT_AI_MONTHLY_BUDGET_USD en .env con el monto mensual en USD aprobado por Administración; configurar PRICEBOT_AI_BUDGET_ALERT_PERCENT (80 por defecto).")
add_number(doc, "Al alcanzar el umbral, PriceBot envía un correo mensual a compras@dynamicenergy.com.ar. El aviso por crédito insuficiente o agotado se envía al detectarlo, con enfriamiento de seis horas.")
add_number(doc, "Consultar saldo y uso en Claude Platform > Billing/Usage. Comprar créditos o habilitar auto reload únicamente con aprobación y medio de pago organizacional.")
add_number(doc, "Una clave API no es saldo: rotar ANTHROPIC_API_KEY no compra créditos y comprar créditos no cambia la clave.")
doc.add_paragraph(
    "El umbral es una estimación calculada con tokens y tarifas de costs_log.jsonl. No consulta el saldo real ni frena el consumo. "
    "Anthropic no ofrece un endpoint de saldo prepago a esta integración. Presupuesto 0 desactiva el umbral, no la alerta por rechazo de saldo."
)

doc.add_heading("Archivos de configuración y secretos", level=1)
add_table(
    doc,
    ["Elemento", "Ubicación", "Respaldo", "Regla"],
    [
        ["Variables y secretos", "pricebot\\.env", "Gestor seguro o copia cifrada", "Nunca subir a Git"],
        ["Usuarios", "pricebot\\api\\data\\users.db", "Copia cifrada periódica", "Detener servicio antes de copiar"],
        ["Código", "Repositorio GitHub", "Git", "No incluir documentos ni secretos"],
        ["Historial de extracción", "Desactivado por defecto", "No aplica", "Activar solo con política de retención"],
        ["Costos", "costs_log.jsonl", "Según política contable", "No registrar nombres de archivo"],
    ],
    [1.25, 2.15, 1.45, 1.85],
)

doc.add_heading("Actualización del sistema", level=1)
reset_numbering()
add_number(doc, "Confirmar que no haya una extracción en curso y cerrar PriceBot.")
add_number(doc, "Respaldar pricebot\\.env y pricebot\\api\\data\\users.db en almacenamiento cifrado.")
add_number(doc, "Ejecutar git status. No continuar si aparecen cambios no reconocidos.")
add_number(doc, "Ejecutar git pull desde la carpeta de instalación indicada por TI.")
add_number(doc, "Si cambiaron dependencias, ejecutar instalar.bat.")
add_number(doc, "Iniciar con iniciar-servidor.bat y completar la lista de verificación posterior.")

doc.add_heading("Controles de seguridad implementados", level=1)
add_table(
    doc,
    ["Riesgo", "Control actual", "Verificación administrativa"],
    [
        ["Acceso no autorizado", "Usuarios confirmados, sesiones firmadas, límite de intentos y administración solo local", "Revisar usuarios mensualmente"],
        ["Sesiones antiguas", "Cambio de contraseña, estado o eliminación revoca sesiones", "Probar login después de cada cambio"],
        ["Archivo falso o malicioso", "Extensiones permitidas, firmas, ZIP Office, rutas, cifrado, tamaño, páginas y píxeles", "Rechazar archivos inesperados"],
        ["Fórmulas al exportar", "Neutralización de fórmulas en XLS y XLSX", "Abrir muestra en entorno controlado"],
        ["Persistencia de información", "Historial, nombres y caché desactivados por defecto", "Verificar .env después de actualizaciones"],
        ["Salida a terceros", "IA externa desactivada por defecto", "Autorizar por clasificación de datos"],
        ["Exposición web", "CORS, validación de origen, CSP y cabeceras de seguridad", "No ampliar orígenes innecesariamente"],
    ],
    [1.45, 3.25, 1.95],
)

doc.add_heading("Riesgos pendientes y mejoras recomendadas", level=1)
add_bullet(doc, "Prioridad alta: instalar y probar el certificado TLS confiable para frontend y API; hasta entonces, el acceso desde otras PCs permanece bloqueado.")
add_bullet(doc, "Configurar respaldo cifrado de users.db y una custodia formal de .env.")
add_bullet(doc, "Asignar propietario y suplente para Claude Platform, Gmail SMTP, GitHub y la PC servidor.")
add_bullet(doc, "Definir política de retención, clasificación y eliminación de documentos procesados.")
add_bullet(doc, "Revisar trimestralmente usuarios, claves, créditos, firewall, actualizaciones y registros de costos.")

doc.add_heading("Respuesta ante incidentes", level=1)
reset_numbering()
add_number(doc, "Detener PriceBot y desconectar el servidor de la red si hay acceso no autorizado o fuga de documentos.")
add_number(doc, "Deshabilitar o eliminar los usuarios afectados.")
add_number(doc, "Rotar ANTHROPIC_API_KEY, contraseña de aplicación Gmail y PRICEBOT_SESSION_SECRET cuando corresponda.")
add_number(doc, "Revisar actividad de Claude Platform, Google, GitHub y Windows.")
add_number(doc, "Preservar registros necesarios para investigación sin copiar documentos confidenciales a medios no aprobados.")
add_number(doc, "Documentar fecha, alcance, archivos involucrados, acciones y responsable del cierre.")

doc.add_heading("Lista de verificación posterior a una actualización", level=1)
checks = [
    "El lanzador inicia backend y frontend sin errores.",
    "La URL local entra como Administrador local sin solicitar contraseña.",
    "La URL LAN solicita correo y contraseña.",
    "Crear, modificar, deshabilitar, habilitar y eliminar usuarios funciona.",
    "El correo de confirmación llega y el enlace usa 192.168.190.146, no 127.0.0.1.",
    "Gmail envía confirmaciones y las alertas llegan a compras@dynamicenergy.com.ar.",
    "HTTP remoto se rechaza y la URL LAN funciona con TLS confiable.",
    "Un archivo pequeño válido se procesa y descarga correctamente.",
    "Un archivo con extensión falsa o tamaño excesivo es rechazado.",
    "La salida coincide con una muestra del documento original.",
    "IA externa solo con saldo aprobado; Git sin documentos ni secretos.",
]
for check in checks:
    add_bullet(doc, "☐ " + check)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_p.add_run("Manual de administración y operación de PriceBot")
footer_run.font.size = Pt(8.5)
footer_run.font.color.rgb = RGBColor(90, 90, 90)

doc.core_properties.title = "Manual de administración y operación de PriceBot"
doc.core_properties.subject = "Operación, usuarios, accesos, tokens, créditos y seguridad"
doc.core_properties.author = "Área de Administración"
doc.core_properties.keywords = "PriceBot administración seguridad usuarios Anthropic"

doc.save(OUTPUT)
print(OUTPUT)
